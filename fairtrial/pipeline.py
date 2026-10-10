"""End-to-end module-B pipeline: documents (from teammate A) -> case cards JSON (for teammate C).

    python -m fairtrial.pipeline --docs data/sample/sample_docs.jsonl --court data/sample/court_records_manual.json --out out/cards.json
    FAIRTRIAL_MOCK=1 python -m fairtrial.pipeline ...   # offline, no API key
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from .agent import CourtRecordStore, follow_up
from .emblematic import precedent_reason, rule_reasons
from .extract import extract_cases
from .llm import LLM
from .resolve import _compatible, matches_known, names_match, resolve, same_province
from .schemas import (
    CaseCard,
    CourtRecordStatus,
    EvidenceClaim,
    QuoteMatch,
    ScreenLabel,
    SourceDocument,
    SourceType,
    SupportStatus,
)
from .screen import screen_with_fallback
from .verify import apply_verified_concerns, evidence_strength, verify_claims

# case fields an agent-found document may fill, only when still unknown and only from a supported claim
FILLABLE = ("province", "court", "case_number", "outlet_name", "charges_original", "charges_en",
            "legal_provisions_original", "detention_status", "work_link",
            "complainant_type", "dewan_pers_status", "outlet_type", "procedural_status")


def load_docs(path: str | Path) -> list[SourceDocument]:
    docs = []
    for line in Path(path).read_text(encoding="utf-8").splitlines():
        if line.strip():
            d = SourceDocument(**json.loads(line))
            d.ensure_paragraphs()
            docs.append(d)
    return docs


def court_claims(case_id: str, rec: dict) -> list[EvidenceClaim]:
    """Structured court data becomes 'court'-type claims (the strongest source grade)."""
    src = rec.get("sipp_url") or f"court:{rec.get('court_name')}"
    rows = [
        ("court", rec.get("court_name")),
        ("case_number", rec.get("case_number")),
        ("legal_provisions_original", rec.get("legal_provisions_original")),
        ("procedural_status", rec.get("procedural_status")),
    ]
    hearings = rec.get("hearing_schedule") or []
    if hearings:
        rows.append(("date_next_hearing", hearings[-1].get("date")))
    out = []
    for field, value in rows:
        if value:
            out.append(EvidenceClaim(case_id=case_id, field=field, value=str(value), source_id=src,
                                     source_type=SourceType.court, paragraph_id="court_record",
                                     exact_quote_original=str(value), quote_match=QuoteMatch.exact,
                                     support_status=SupportStatus.supported,
                                     verifier_rationale=f"manual SIPP lookup ({rec.get('retrieved_at')}), see screenshot"))
    return out


def _same_person(case, other) -> bool:
    """Same rules as resolve(): a variant/short name match also needs the same province."""
    for name in [case.person_name_original, *case.person_aliases]:
        ok, rule = names_match(name, other.person_name_original)
        if ok and _compatible(case, other) and (rule == "exact" or same_province(case.province, other.province)):
            return True
    return False


def run(docs: list[SourceDocument], llm: LLM, court_store: CourtRecordStore, known_international: set[str] | None = None,
        search_factory=None, coverage: bool = False, partner_records: list[dict] | None = None,
        partner_reference: list[dict] | None = None):
    """search_factory(case, log_list, max_docs) -> search_fn(query) -> list[SourceDocument]; None = no network search.
    coverage=True adds the media-coverage step (fairtrial/coverage.py) to every card; it searches only if
    search_factory is set, otherwise it matches against this run's own documents."""
    log = {"screening": [], "agent": [], "duplicate_candidates": [], "field_conflicts": [],
           "failed_documents": [], "failed_cases": [], "court_records": [],
           "agent_search": [], "agent_documents": [], "coverage_search": [], "coverage": []}
    by_id = {d.document_id: d for d in docs}
    extracted = []

    def _failed(doc_id: str, step: str, exc: Exception, **extra) -> None:
        """One bad document must not take the whole run down with it."""
        log["failed_documents"].append({"document_id": doc_id, "step": step,
                                        "error": f"{type(exc).__name__}: {exc}", **extra})
        print(f"!! {doc_id} failed at {step}: {type(exc).__name__}: {exc}")

    for d in docs:
        try:
            s = screen_with_fallback(d, llm)
        except Exception as e:
            _failed(d.document_id, "screening", e)
            continue
        log["screening"].append(s.model_dump())
        if s.label == ScreenLabel.irrelevant:
            continue  # metadata kept by teammate A; sample of these is re-checked for misses
        try:
            cases = extract_cases(d, llm)
        except Exception as e:
            _failed(d.document_id, "extraction", e)
            continue
        for case, claims in cases:
            try:
                verify_claims(claims, d, llm)
            except Exception as e:
                # claims keep their default needs_review status, so nothing unverified
                # can be presented as supported
                _failed(d.document_id, "verification", e, case_id=case.case_id)
            extracted.append((case, claims))

    merged, dups, field_conflicts = resolve(extracted)
    log["duplicate_candidates"] = dups
    log["field_conflicts"] = field_conflicts

    def _failed_case(case_id: str, step: str, exc: Exception) -> None:
        """A card is still worth producing without its agent step or its reasons."""
        log["failed_cases"].append({"case_id": case_id, "step": step,
                                    "error": f"{type(exc).__name__}: {exc}"})
        print(f"!! {case_id} failed at {step}: {type(exc).__name__}: {exc}")

    def _absorb_document(case, doc: SourceDocument) -> list[EvidenceClaim]:
        """Agent-found document -> screen -> extract -> verify -> claims for THIS case only.
        Nothing reaches the card without going through verify_claims, so an unverified claim
        can never show up as supported."""
        doc.ensure_paragraphs()
        entry = {"case_id": case.case_id, "document_id": doc.document_id, "url": doc.source_url}
        log["agent_documents"].append(entry)
        if doc.document_id in case.source_ids or any(d.source_url == doc.source_url for d in by_id.values()):
            entry["result"] = "already a source in this run"
            return []
        s = screen_with_fallback(doc, llm)
        entry["screen"] = s.label.value
        if s.label == ScreenLabel.irrelevant:
            entry["result"] = "screened irrelevant"
            return []
        found = extract_cases(doc, llm)
        mine = [(c, cl) for c, cl in found if _same_person(case, c)]
        entry["other_people_ignored"] = [c.person_name_original for c, _ in found if not _same_person(case, c)]
        if not mine:
            entry["result"] = "person not extracted from document"
            return []
        by_id[doc.document_id] = doc          # so precedent_reason can see its paragraphs
        case.source_ids = sorted(set(case.source_ids) | {doc.document_id})
        added: list[EvidenceClaim] = []
        for other, cl in mine:
            verify_claims(cl, doc, llm)
            for c in cl:
                c.case_id = case.case_id
                if c.support_status != SupportStatus.supported:
                    continue
                if c.field in FILLABLE:
                    cur = getattr(case, c.field)
                    cur = cur.value if hasattr(cur, "value") else cur
                    new = getattr(other, c.field)
                    new = new.value if hasattr(new, "value") else new
                    if cur == "unknown":
                        setattr(case, c.field, new)
                    elif str(cur).lower() != str(new).lower():
                        # disagreement goes to a human, same policy as resolve()
                        c.support_status = SupportStatus.needs_review
                        log["field_conflicts"].append({"case_id": case.case_id, "field": c.field, "kept": cur,
                                                       "values": [{"value": cur, "source_ids": [s for s in case.source_ids if s != doc.document_id]},
                                                                  {"value": new, "source_ids": [doc.document_id]}]})
                elif c.field.startswith("date_"):
                    case.event_dates.setdefault(c.field.removeprefix("date_"), other.event_dates.get(c.field.removeprefix("date_"), c.value))
            added += cl
        entry["result"] = f"{len(added)} claims, {sum(c.support_status == SupportStatus.supported for c in added)} supported"
        return added

    from collections import Counter
    doc_case_count = Counter(sid for c, _ in merged for sid in c.source_ids)   # >1 = a round-up article

    cards = []
    for case, claims in merged:
        try:
            search_fn = search_factory(case, log["agent_search"], 3) if search_factory else None
            case, new_evidence, trace = follow_up(case, claims, llm, court_store, search_fn=search_fn)
            log["agent"].append({"case_id": case.case_id, "steps": trace.steps})
        except Exception as e:
            _failed_case(case.case_id, "follow_up", e)
            new_evidence = []
        try:
            for ev in new_evidence:
                if ev["kind"] == "court_record":
                    rec = ev["record"]
                    log["court_records"].append({"case_id": case.case_id, "record": rec})
                    claims += court_claims(case.case_id, rec)
                    for f in ("court", "case_number", "legal_provisions_original"):
                        if getattr(case, f) == "unknown" and rec.get(f if f != "court" else "court_name"):
                            setattr(case, f, rec.get(f if f != "court" else "court_name"))
                    if rec.get("hearing_schedule"):
                        case.event_dates.setdefault("next_hearing", rec["hearing_schedule"][-1]["date"])
        except Exception as e:
            _failed_case(case.case_id, "court_record", e)
        for ev in new_evidence:
            if ev["kind"] != "document":
                continue
            try:
                claims += _absorb_document(case, ev["doc"])
            except Exception as e:
                _failed_case(case.case_id, "agent_document", e)

        try:
            apply_verified_concerns(case, claims)
            if known_international is not None:
                case.known_internationally = matches_known(case, known_international)
            reasons = rule_reasons(case, claims)
        except Exception as e:
            _failed_case(case.case_id, "emblematic", e)
            reasons = []
        try:
            paragraphs = {pid: t for sid in case.source_ids for pid, t in by_id[sid].paragraphs.items()}
            pr = precedent_reason(case, claims, paragraphs, llm)
            if pr:
                reasons.append(pr)
        except Exception as e:
            _failed_case(case.case_id, "precedent_reason", e)

        # freedom-of-expression sign, case description, partner involvement: rules only, no model call
        foe = ctx = partners = None
        try:
            from .context import build_context
            from .foe import assess_foe
            from .partners import assess_partners
            from .schemas import EmblematicReason
            foe = assess_foe(case, claims)
            if foe.value in ("Yes", "Possible"):
                refs = sorted({f"{e['source_id']}#{e['paragraph_id']}" for s in foe.signals for e in s.evidence
                               if s.signal in ("link", "speech_law", "press")})
                reasons.append(EmblematicReason(code="freedom_of_expression", method="rule", evidence_refs=refs,
                                                explanation=f"{foe.value}: {foe.rationale}"))
            ctx = build_context(case, claims, foe)
            partners = assess_partners(case, by_id, doc_case_count, partner_records or [], partner_reference)
        except Exception as e:
            _failed_case(case.case_id, "foe_context_partners", e)

        cov = None
        if coverage:
            try:
                from .coverage import build_coverage
                cov_search = search_factory(case, log["coverage_search"], 10) if search_factory else None
                cov = build_coverage(case, by_id, llm, search_fn=cov_search, doc_case_count=doc_case_count)
                log["coverage"].append({"case_id": case.case_id, "n_articles": cov.n_articles,
                                        "n_outlets": cov.n_outlets, "n_independent": cov.n_independent,
                                        "n_uncertain": cov.n_uncertain})
            except Exception as e:
                _failed_case(case.case_id, "coverage", e)
        try:
            cards.append(CaseCard(case=case, claims=claims, emblematic_reasons=reasons,
                                  evidence_strength=evidence_strength(claims),
                                  coverage=cov.model_dump() if cov else None,
                                  foe=foe.model_dump() if foe else None,
                                  context=ctx.model_dump() if ctx else None,
                                  partners=partners.model_dump() if partners else None))
        except Exception as e:
            _failed_case(case.case_id, "card", e)
    return cards, log


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--docs", default="data/sample/sample_docs.jsonl")
    ap.add_argument("--court", default="data/sample/court_records_manual.json")
    ap.add_argument("--known", default="data/sample/known_international.txt", help="one lower-case name per line")
    ap.add_argument("--out", default="out/cards.json")
    ap.add_argument("--db", help="read from / write to A's trialwatch.db instead of --docs (never pass gold.db)")
    ap.add_argument("--prompt-version", default="v1")
    ap.add_argument("--search", action="store_true",
                    help="let the follow-up agent search the six RSS sources (network; off by default)")
    ap.add_argument("--partners", default="data/partner_records_manual.json",
                    help="hand-checked partner involvement records (JSON list); the only source of 'Confirmed'")
    ap.add_argument("--partner-reference", default="fairtrial/partners_ref.json",
                    help="organisations to look for in articles. The shipped file is a DEMO list; in production "
                         "point this at TrialWatch's internal partner directory (same JSON format)")
    ap.add_argument("--coverage", action="store_true",
                    help="add media coverage (other reports of the same case) to every card; searches only with --search")
    ap.add_argument("--contact", default="", help="real e-mail for the User-Agent, required with --search")
    ap.add_argument("--exclude-workbook", default="data/workbooks/TrialWatch_Indonesia_Annotation_Workbook.xlsx",
                    help="with --search: never return an article listed in this workbook (any split)")
    args = ap.parse_args()
    if args.search and not args.contact:
        raise SystemExit("--search needs --contact you@example.com")

    from .cache import make_llm
    llm = make_llm()          # FAIRTRIAL_LLM_CACHE=<path> -> identical calls are not paid twice
    if llm.mock:
        from .mock import register_all
        register_all(llm)

    known = None
    if Path(args.known).exists():
        known = {l.strip().lower() for l in Path(args.known).read_text(encoding="utf-8").splitlines()
                 if l.strip() and not l.startswith("#")}

    if args.db:
        if "gold" in Path(args.db).name.lower():
            raise SystemExit("refusing to run on a gold database")
        from .db_io import load_docs_from_db
        docs = load_docs_from_db(args.db)
    else:
        docs = load_docs(args.docs)
    search_factory = None
    if args.search:
        from .search import CachingFetcher, make_rss_search, norm_url
        fetcher = CachingFetcher(args.contact)
        exclude = set()
        if Path(args.exclude_workbook).exists():
            from openpyxl import load_workbook
            ws = load_workbook(args.exclude_workbook, read_only=True, data_only=True)["Articles"]
            rows = ws.iter_rows(values_only=True)
            col = list(next(rows)).index("URL")
            exclude = {norm_url(r[col]) for r in rows if r[col]}
            print(f"agent search: excluding {len(exclude)} labelled article URLs")
        else:
            print(f"!! {args.exclude_workbook} not found: agent search may return labelled articles")
        search_factory = lambda case, lg, max_docs=3: make_rss_search(
            args.contact, [case.person_name_original, *case.person_aliases], exclude_urls=exclude, log=lg,
            max_docs=max_docs, fetcher=fetcher)
    from .partners import load_manual, load_reference
    cards, log = run(docs, llm, CourtRecordStore(args.court), known, search_factory=search_factory,
                     coverage=args.coverage, partner_records=load_manual(args.partners),
                     partner_reference=load_reference(args.partner_reference))

    if args.db:
        from .db_io import write_court_record, write_results
        from .llm import FAST_MODEL
        counts = write_results(args.db, cards, log["screening"], model="mock" if llm.mock else FAST_MODEL,
                               prompt_version=args.prompt_version)
        for cr in log["court_records"]:
            write_court_record(args.db, cr["case_id"], cr["record"])
        from .card_db import write_case_context
        print(f"wrote {write_case_context(args.db, cards)} rows to case_context (FOE, description, partners, coverage)")
        if args.coverage:
            from .coverage import Coverage, write_coverage
            n_cov = sum(write_coverage(args.db, c.case.case_id, Coverage(**c.coverage)) for c in cards if c.coverage)
            print(f"wrote {n_cov} coverage rows to case_coverage")
        print(f"wrote to {args.db}: {counts}, court_records={len(log['court_records'])}")
    Path(args.out).parent.mkdir(parents=True, exist_ok=True)
    Path(args.out).write_text(json.dumps([c.model_dump(mode="json") for c in cards], ensure_ascii=False, indent=2), encoding="utf-8")
    Path(args.out).with_name("run_log.json").write_text(json.dumps(log, ensure_ascii=False, indent=2, default=str), encoding="utf-8")

    for c in cards:
        n_ok = sum(x.support_status == SupportStatus.supported for x in c.claims)
        n_bad = sum(x.quote_match == QuoteMatch.not_found for x in c.claims)
        print(f"{c.case.person_name_original:<20} stage={c.case.procedural_status.value:<12} "
              f"strength={c.evidence_strength:<11} claims={len(c.claims):>2} supported={n_ok:>2} "
              f"quote_not_found={n_bad} court={c.case.court_record_status.value} "
              f"emblematic={[r.code for r in c.emblematic_reasons]}"
              + (f" foe={c.foe['value']}" if c.foe else "")
              + (f" partners={c.partners['status']}" if c.partners else "")
              + (f" coverage={c.coverage['n_articles']} articles/{c.coverage['n_outlets']} outlets/"
                 f"{c.coverage['n_independent']} independent" if c.coverage else ""))
    print(f"wrote {args.out}")
    if hasattr(llm, "hits"):
        print(f"model cache: {llm.hits} hits (free), {llm.misses} new calls")


if __name__ == "__main__":
    main()
