"""Record that the answer key was reviewed: set "Review status" to Reviewed in the annotation workbook.

    python scripts/mark_reviewed.py --reviewer "Name(s)" --sheets Cases,Claims,Pairs
    python scripts/mark_reviewed.py --reviewer "..." --sheets Cases --ids C001,C002   # only some rows

Use it only after people have actually checked the rows. It changes nothing but the status, and writes
who reviewed and when into the row's Notes column (if the sheet has one) and into a "Review log" sheet,
so the workbook itself shows the review was recorded, by whom and on what date.
A backup copy of the workbook is written next to it first.
"""
from __future__ import annotations

import argparse
import shutil
import sys
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

WB = "data/workbooks/TrialWatch_Indonesia_Annotation_Workbook.xlsx"


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", default=WB)
    ap.add_argument("--reviewer", required=True, help="who reviewed (recorded in the workbook)")
    ap.add_argument("--sheets", default="Cases", help="comma-separated sheet names, e.g. Cases,Claims,Pairs")
    ap.add_argument("--ids", default="", help="optional comma-separated row ids (first column) to limit to")
    ap.add_argument("--status", default="Reviewed")
    a = ap.parse_args()

    path = Path(a.workbook)
    backup = path.with_name(path.stem + f".before_review_{date.today().isoformat()}.xlsx")
    shutil.copy2(path, backup)
    wb = load_workbook(path)
    only = {x.strip() for x in a.ids.split(",") if x.strip()}
    today = date.today().isoformat()
    log = []
    for name in [s.strip() for s in a.sheets.split(",") if s.strip()]:
        ws = wb[name]
        header = [c.value for c in ws[1]]
        if "Review status" not in header:
            print(f"!! {name}: no 'Review status' column, skipped")
            continue
        col = header.index("Review status") + 1
        note_col = next((i + 1 for i, h in enumerate(header) if h and str(h).strip().lower() in ("notes", "note")), None)
        n = 0
        for row in ws.iter_rows(min_row=2):
            rid = row[0].value
            status = row[col - 1].value
            if not rid or status in (None, "Example") or str(rid).startswith(("A000", "C000", "example")):
                continue
            if only and str(rid) not in only:
                continue
            if status != a.status:
                row[col - 1].value = a.status
                if note_col:
                    old = row[note_col - 1].value or ""
                    row[note_col - 1].value = (f"{old} | " if old else "") + f"Reviewed by {a.reviewer} on {today}"
                n += 1
        log.append((name, n))
        print(f"{name}: {n} row(s) set to {a.status}")
    ws = wb["Review log"] if "Review log" in wb.sheetnames else wb.create_sheet("Review log")
    if ws.max_row == 1 and ws.cell(1, 1).value is None:
        ws.append(["Date", "Reviewer", "Sheet", "Rows set to " + a.status, "Ids"])
    for name, n in log:
        ws.append([today, a.reviewer, name, n, a.ids or "all"])
    wb.save(path)
    print(f"saved {path} (backup: {backup.name})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
