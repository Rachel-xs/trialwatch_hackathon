#!/usr/bin/env python3
"""
Batch-download the articles listed in the Articles sheet of the annotation workbook.

What it does, per article:
  1. Uses a manually saved file if one exists (data/manual/A001.txt). This is the fallback for anything blocked.
  2. Checks robots.txt. If the page is disallowed it is NOT fetched; it is logged as blocked_robots for manual saving.
  3. Downloads the page politely (one request per domain every few seconds, identified User-Agent).
  4. Extracts the main text (HTML via trafilatura, PDF via pypdf), detects the language, hashes the text.
  5. Writes data/raw, data/text, data/meta and appends to data/fetch_log.csv.

It never tries to get around a block, a login or a paywall.

Usage:
  python fetch_articles.py --workbook TrialWatch_Annotation_Workbook.xlsx --contact you@example.com
  python fetch_articles.py --workbook ... --ids A001,A002      # only some articles
  python fetch_articles.py --workbook ... --split Tuning        # only the tuning set
  python fetch_articles.py --workbook ... --force               # re-download even if cached
"""
import argparse, csv, hashlib, io, json, re, sys, time
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse
from urllib import robotparser

import requests
from openpyxl import load_workbook

PAYWALL_HINTS = ["subscribe", "subscription", "paywall", "abone ol", "üye ol", "üyelik", "berlangganan", "log in to read", "sign in to read"]
MIN_FULL_CHARS = 400
LOG_FIELDS = ["article_id", "url", "final_url", "http_status", "fetch_status", "fetch_source", "chars",
              "language_detected", "published_extracted", "retrieved_at", "note"]


def read_articles(workbook):
    ws = load_workbook(workbook, data_only=True)["Articles"]
    headers = [c.value for c in ws[1]]
    rows = []
    for r in ws.iter_rows(min_row=3, values_only=True):          # row 2 is the example row
        rec = dict(zip(headers, r))
        if rec.get("Article ID"):
            rows.append(rec)
    return rows


def now():
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


class Fetcher:
    def __init__(self, contact, delay):
        self.s = requests.Session()
        self.s.headers.update({"User-Agent": f"TrialWatchHackathonBot/0.1 (academic research; contact: {contact})",
                               "Accept-Language": "en,tr;q=0.8,id;q=0.8"})
        self.delay = delay
        self.last = {}
        self.robots = {}

    def _wait(self, domain):
        gap = time.time() - self.last.get(domain, 0)
        if gap < self.delay:
            time.sleep(self.delay - gap)
        self.last[domain] = time.time()

    def allowed(self, url):
        p = urlparse(url)
        base = f"{p.scheme}://{p.netloc}"
        if base not in self.robots:
            rp = robotparser.RobotFileParser()
            try:
                self._wait(p.netloc)
                r = self.s.get(base + "/robots.txt", timeout=15)
                if r.status_code == 200:
                    rp.parse(r.text.splitlines())
                elif r.status_code in (401, 403):
                    rp.parse(["User-agent: *", "Disallow: /"])
                else:
                    rp.parse([])                                  # no robots.txt: allowed
            except requests.RequestException:
                rp.parse([])
            self.robots[base] = rp
        return self.robots[base].can_fetch(self.s.headers["User-Agent"], url)

    def get(self, url, tries=2):
        domain = urlparse(url).netloc
        last_exc = None
        for i in range(tries):
            self._wait(domain)
            try:
                r = self.s.get(url, timeout=30, allow_redirects=True)
                if r.status_code in (429, 503) and i < tries - 1:
                    time.sleep(10 * (i + 1))
                    continue
                return r
            except requests.RequestException as e:
                last_exc = e
                time.sleep(3)
        raise last_exc


def extract_html(html, url):
    title, text, date = None, "", None
    try:
        import trafilatura
        text = trafilatura.extract(html, url=url, include_comments=False, include_tables=False, favor_recall=True) or ""
        meta = trafilatura.extract_metadata(html)
        if meta:
            title, date = meta.title, meta.date
    except Exception:
        pass
    if len(text) < MIN_FULL_CHARS:                                # fallback
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(html, "html.parser")
        for t in soup(["script", "style", "nav", "footer", "header", "aside", "form"]):
            t.decompose()
        root = soup.find("article") or soup.find("main") or soup.body or soup
        paras = [p.get_text(" ", strip=True) for p in root.find_all("p")]
        alt = "\n\n".join(p for p in paras if len(p) > 40)
        if len(alt) > len(text):
            text = alt
        if not title and soup.title:
            title = soup.title.get_text(strip=True)
    return title, text.strip(), date


def extract_pdf(data):
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join((p.extract_text() or "").strip() for p in reader.pages).strip()


def detect_lang(text):
    try:
        from langdetect import DetectorFactory, detect
        DetectorFactory.seed = 0
        return detect(text[:3000]) if text else None
    except Exception:
        return None


def save(out, aid, raw, raw_ext, text, meta):
    (out / "raw").mkdir(parents=True, exist_ok=True)
    (out / "text").mkdir(parents=True, exist_ok=True)
    (out / "meta").mkdir(parents=True, exist_ok=True)
    if raw is not None:
        (out / "raw" / f"{aid}.{raw_ext}").write_bytes(raw)
    (out / "text" / f"{aid}.txt").write_text(text, encoding="utf-8")
    (out / "meta" / f"{aid}.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")


def process(rec, f, out, force):
    aid, url = rec["Article ID"], rec["URL"]
    base = {"article_id": aid, "url": url, "final_url": "", "http_status": "", "fetch_source": "web",
            "chars": 0, "language_detected": "", "published_extracted": "", "retrieved_at": now(), "note": ""}

    manual = out / "manual" / f"{aid}.txt"
    if manual.exists():
        text = manual.read_text(encoding="utf-8").strip()
        meta = {**base, "fetch_status": "full" if len(text) >= MIN_FULL_CHARS else "title_only", "fetch_source": "manual",
                "title": rec.get("Title (original)"), "language": detect_lang(text), "language_detected": detect_lang(text) or "", "published_extracted": None,
                "content_hash": hashlib.sha256(text.encode()).hexdigest()}
        save(out, aid, None, None, text, meta)
        return {**base, "fetch_status": meta["fetch_status"], "fetch_source": "manual", "chars": len(text),
                "language_detected": meta["language"] or "", "note": "from data/manual"}

    if (out / "meta" / f"{aid}.json").exists() and not force:
        m = json.loads((out / "meta" / f"{aid}.json").read_text(encoding="utf-8"))
        if m.get("fetch_status") in ("full", "title_only", "paywall"):
            return {**base, "fetch_status": m["fetch_status"], "fetch_source": m.get("fetch_source", "web"),
                    "chars": m.get("chars", 0), "note": "cached"}

    if not f.allowed(url):
        return {**base, "fetch_status": "blocked_robots", "note": "robots.txt disallows; save the text by hand to data/manual/"}
    try:
        r = f.get(url)
    except requests.RequestException as e:
        return {**base, "fetch_status": "failed", "note": f"network error: {type(e).__name__}"}
    base.update(final_url=r.url, http_status=r.status_code)
    if r.status_code != 200:
        return {**base, "fetch_status": "failed", "note": f"HTTP {r.status_code}; save the text by hand to data/manual/"}

    ctype = r.headers.get("Content-Type", "").lower()
    if "pdf" in ctype or url.lower().endswith(".pdf"):
        text, title, date, raw_ext = extract_pdf(r.content), rec.get("Title (original)"), None, "pdf"
        html_for_hint = ""
    else:
        r.encoding = r.apparent_encoding if (r.encoding in (None, "ISO-8859-1")) else r.encoding
        html = r.text
        title, text, date = extract_html(html, r.url)
        raw_ext, html_for_hint = "html", html.lower()

    if len(text) >= MIN_FULL_CHARS:
        status = "full"
    elif any(h in html_for_hint for h in PAYWALL_HINTS):
        status = "paywall"
    else:
        status = "title_only"
    lang = detect_lang(text)
    meta = {**base, "fetch_status": status, "title": title or rec.get("Title (original)"), "language": lang, "language_detected": lang or "",
            "published_extracted": str(date) if date else None, "chars": len(text),
            "content_hash": hashlib.sha256(text.encode()).hexdigest()}
    save(out, aid, r.content, raw_ext, text, meta)
    note = "" if status == "full" else "text too short; save the text by hand to data/manual/"
    return {**base, "fetch_status": status, "chars": len(text), "language_detected": lang or "",
            "published_extracted": str(date) if date else "", "note": note}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", required=True)
    ap.add_argument("--out", default="data")
    ap.add_argument("--contact", default="team@example.com", help="your email, shown in the User-Agent")
    ap.add_argument("--delay", type=float, default=4.0, help="seconds between requests to the same domain")
    ap.add_argument("--ids", default="", help="comma-separated Article IDs")
    ap.add_argument("--skip", default="", help="comma-separated Article IDs to leave out entirely")
    ap.add_argument("--split", default="", choices=["", "Tuning", "Final test"])
    ap.add_argument("--force", action="store_true")
    a = ap.parse_args()

    rows = read_articles(a.workbook)
    if a.ids:
        want = {x.strip() for x in a.ids.split(",")}
        rows = [r for r in rows if r["Article ID"] in want]
    if a.split:
        rows = [r for r in rows if r.get("Data split") == a.split]
    if a.skip:
        drop = {x.strip() for x in a.skip.split(",") if x.strip()}
        present = sorted(drop & {r["Article ID"] for r in rows})
        rows = [r for r in rows if r["Article ID"] not in drop]
        if present:
            print(f"skipping {', '.join(present)} (--skip): not fetched, not written to the log")
    out = Path(a.out); out.mkdir(parents=True, exist_ok=True)
    f = Fetcher(a.contact, a.delay)

    results = []
    for i, rec in enumerate(rows, 1):
        res = process(rec, f, out, a.force)
        results.append(res)
        print(f"[{i}/{len(rows)}] {res['article_id']}  {res['fetch_status']:<14} {res['chars']:>6} chars  {res['note']}")

    log = out / "fetch_log.csv"
    new = not log.exists()
    with log.open("a", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=LOG_FIELDS)
        if new:
            w.writeheader()
        for res in results:
            w.writerow({k: res.get(k, "") for k in LOG_FIELDS})

    counts = {}
    for res in results:
        counts[res["fetch_status"]] = counts.get(res["fetch_status"], 0) + 1
    print("\nSummary:", counts)
    todo = [r["article_id"] for r in results if r["fetch_status"] != "full"]
    if todo:
        print("Save these by hand (copy the article text into data/manual/<ID>.txt), then run this script again:")
        print("  " + ", ".join(todo))


if __name__ == "__main__":
    sys.exit(main())
