#!/usr/bin/env python3
"""Discover ARTICLE URLs from 6 RSS/Atom/sitemap sources; write workbook compatible with fetch_articles.py.
Does not bypass robots.txt, paywalls, logins or access restrictions. Requires requests, openpyxl.
"""
import argparse, datetime as dt, hashlib, re, time, xml.etree.ElementTree as ET
from urllib.parse import urlparse
from urllib.robotparser import RobotFileParser
import requests
from openpyxl import Workbook, load_workbook

HEADERS=['Article ID','Owner','Country','Language','Publisher','URL','Title (original)','Published','Retrieved','Fetch status','Relevance','Test type','Case ID','Notes','Review status','Data split']
KEYWORDS=['wartawan','jurnalis','jurnalistik','pers','journalist','reporter','ditangkap','ditahan','tersangka','pengadilan','dakwaan','pidana','pencemaran nama baik','press freedom','kebebasan pers']
NS_ATOM='{http://www.w3.org/2005/Atom}'

def norm_date(s):
    if not s: return 'TBD'
    try:
        from email.utils import parsedate_to_datetime
        return parsedate_to_datetime(s).date().isoformat()
    except Exception: pass
    m=re.search(r'\d{4}-\d{2}-\d{2}',s)
    return m.group(0) if m else 'TBD'

def allowed(session,url):
    p=urlparse(url); root=f'{p.scheme}://{p.netloc}'
    try:
        r=session.get(root+'/robots.txt',timeout=12)
        if r.status_code in (401,403): return False
        if r.status_code==200:
            rp=RobotFileParser();rp.parse(r.text.splitlines())
            return rp.can_fetch(session.headers['User-Agent'],url)
        return True
    except requests.RequestException:
        print('  robots.txt unavailable: skipping feed for conservative access policy')
        return False

def get_text(elem,path):
    x=elem.find(path)
    return (x.text or '').strip() if x is not None else ''

def parse_feed(content,feed_url):
    root=ET.fromstring(content)
    out=[]
    if root.tag.endswith('rss') or root.tag.endswith('RDF'):
        for item in root.findall('.//item'):
            out.append((get_text(item,'link'),get_text(item,'title'),get_text(item,'pubDate')))
    elif root.tag==NS_ATOM+'feed' or root.tag.endswith('feed'):
        for item in root.findall('.//'+NS_ATOM+'entry'):
            link=''
            for a in item.findall(NS_ATOM+'link'):
                if a.attrib.get('rel','alternate')=='alternate' and a.attrib.get('href'):
                    link=a.attrib['href'];break
            out.append((link,get_text(item,NS_ATOM+'title'),get_text(item,NS_ATOM+'published') or get_text(item,NS_ATOM+'updated')))
    elif root.tag.endswith('urlset'):
        for item in root:
            out.append((get_text(item,'{*}loc'),'',''))
    elif root.tag.endswith('sitemapindex'):
        print('  Sitemap index detected: nested sitemaps not followed in MVP.')
    else:
        raise ValueError('Not a recognized RSS/Atom/sitemap XML format')
    return out

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--sources',default='data/workbooks/TrialWatch_Six_Sources_Ready.xlsx')
    ap.add_argument('--output',default='data/workbooks/TrialWatch_Live_Articles.xlsx')
    ap.add_argument('--limit-per-source',type=int,default=15)
    ap.add_argument('--only-keywords',action='store_true',help='Exclude titles without keywords (may miss relevant cases)')
    ap.add_argument('--contact',required=True,help='Real contact email for User-Agent')
    args=ap.parse_args()
    src=load_workbook(args.sources,data_only=True)['Sources']
    session=requests.Session();session.headers['User-Agent']=f'TrialWatchResearchPrototype/0.1 (contact: {args.contact})'
    discovered=[];seen=set();log=[]
    for name,kind,url,*_ in src.iter_rows(min_row=2,values_only=True):
        if not url: continue
        print(f'Checking {name}: {url}')
        if not allowed(session,url):
            print('  Skipped: robots restriction or robots check unavailable');log.append((name,'blocked_or_unverified_robots',0));continue
        try:
            r=session.get(url,timeout=25);r.raise_for_status()
            entries=parse_feed(r.content,url)
            n=0
            for link,title,published in entries:
                if n>=args.limit_per_source:break
                if not link or not urlparse(link).scheme.startswith('http') or link in seen:continue
                hit=any(k in (title or '').lower() for k in KEYWORDS)
                if args.only_keywords and not hit:continue
                seen.add(link);n+=1
                discovered.append((name,link,title or '',norm_date(published),hit))
            log.append((name,'ok',n));print(f'  Found {n} article links')
        except (requests.RequestException,ET.ParseError,ValueError) as exc:
            log.append((name,f'error: {type(exc).__name__}',0));print(f'  Could not parse source: {exc}')
        time.sleep(2)
    wb=Workbook();ws=wb.active;ws.title='Articles';ws.append(HEADERS)
    ws.append(['A000 (example)','', 'Indonesia','id','example','https://example.com/article','EXAMPLE — ignore','TBD','','','uncertain','','','Do not fetch example','Example','Tuning'])
    for i,(name,url,title,published,hit) in enumerate(discovered,1):
        ws.append(['L'+hashlib.sha1(url.encode()).hexdigest()[:8],'','Indonesia','id',name,url,title,published,'','','uncertain','','',f'Feed discovery; keyword title match: {hit}. NOT verified.','Unreviewed','Tuning'])
    ws.freeze_panes='A3';ws.auto_filter.ref=f'A1:P{ws.max_row}'
    for col,width in {'A':15,'E':22,'F':80,'G':70,'H':18,'N':55}.items():ws.column_dimensions[col].width=width
    notes=wb.create_sheet('Discovery_Log');notes.append(['Source','Status','Articles exported'])
    for row in log:notes.append(row)
    wb.save(args.output)
    print(f'\nSaved {len(discovered)} discovered article links to {args.output}')
    print('These are UNVERIFIED article candidates. Check titles/URLs and sample fetch results manually.')
    if not discovered:print('No article URLs found. Check Discovery_Log; do not run fetch_articles.py until URLs are available.')
if __name__=='__main__':main()
